import hashlib
import io
import shutil

import pytest
from docx import Document as DocxDocument

from findr.adapters.outbound.files.mime_types import DOCX, MARKDOWN, PDF, TEXT
from findr.adapters.outbound.files.unstructured_document_parser import (
    CHUNKING_VERSION,
    MAX_CHARACTERS,
    UnstructuredDocumentParser,
)
from findr.domain.exceptions import ExtractionFailed, UnsupportedFileType
from findr.domain.value_objects import ChunkKind

# PDFs go through unstructured's hi_res strategy (layout model, OCR, table
# model), which needs tesseract and poppler on PATH. They're installed in the
# Docker image; elsewhere (`brew install tesseract poppler` on macOS) these
# tests are skipped.
needs_ocr = pytest.mark.skipif(
    not (shutil.which("tesseract") and shutil.which("pdftoppm")),
    reason="needs tesseract and poppler (installed in the Docker image)",
)

ROWS = [["Plan", "Monthly price", "Seats"], ["Starter", "10 GBP", "5"], ["Business", "45 GBP", "50"]]


@pytest.fixture(scope="module")
def parser() -> UnstructuredDocumentParser:
    return UnstructuredDocumentParser()


def _parse(parser, content: bytes, mime_type: str, filename: str = "file", version: int = 1):
    return parser.parse(content, mime_type, filename, version)


def _docx(*, paragraphs=(), table_rows=None) -> bytes:
    doc = DocxDocument()
    for kind, text in paragraphs:
        if kind == "h":
            doc.add_heading(text, 1)
        else:
            doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell in enumerate(row):
                table.rows[r].cells[c].text = cell
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_plain_text_becomes_body_text_and_chunks(parser):
    content = b"Renewal pricing rises next year.\n\nSupport hours are nine to five."
    parsed = _parse(parser, content, TEXT, "notes.txt")

    assert "Renewal pricing rises next year." in parsed.text
    assert "Support hours are nine to five." in parsed.text
    assert parsed.chunks
    assert all(c.kind == ChunkKind.TEXT and c.table_html is None for c in parsed.chunks)


def test_long_text_is_split_into_bounded_overlapping_chunks(parser):
    sentence = "The quarterly renewal pricing for the Acme contract rises twenty percent. "
    parsed = _parse(parser, (sentence * 60).encode(), TEXT)

    assert len(parsed.chunks) > 1
    assert all(len(c.text) <= MAX_CHARACTERS for c in parsed.chunks)
    # The split pieces of one long element overlap and are flagged.
    assert parsed.chunks[1].text[:40] in parsed.chunks[0].text
    assert parsed.chunks[1].metadata.is_continuation
    assert not parsed.chunks[0].metadata.is_continuation


def test_markdown_sections_and_tables(parser):
    content = (
        "# Pricing\n\nThe subscription cost has been updated.\n\n"
        "| Plan | Monthly price | Seats |\n|---|---|---|\n| Starter | 10 GBP | 5 |\n"
        "| Business | 45 GBP | 50 |\n\n"
        "# Support\n\nSupport hours are nine to five on weekdays.\n"
    ).encode()
    parsed = _parse(parser, content, MARKDOWN, "pricing.md")

    tables = [c for c in parsed.chunks if c.kind == ChunkKind.TABLE]
    assert len(tables) == 1
    assert "Starter" in tables[0].text and "45 GBP" in tables[0].text
    # Chunking normalises header cells to <td>; rows and columns are intact.
    assert "<td>Monthly price</td>" in tables[0].table_html
    assert "<tr><td>Business</td><td>45 GBP</td><td>50</td></tr>" in tables[0].table_html
    # A table is never merged with prose.
    assert "subscription" not in tables[0].text
    sections = {c.metadata.section_title for c in parsed.chunks}
    assert {"Pricing", "Support"} <= sections
    support = next(c for c in parsed.chunks if "Support hours" in c.text)
    assert support.metadata.section_title == "Support"


def test_docx_table_is_kept_exactly(parser):
    content = _docx(
        paragraphs=[("h", "Pricing"), ("p", "The subscription cost has been updated.")],
        table_rows=ROWS,
    )
    parsed = _parse(parser, content, DOCX, "pricing.docx")

    [table] = [c for c in parsed.chunks if c.kind == ChunkKind.TABLE]
    for row in ROWS:
        for cell in row:
            assert f"<td>{cell}</td>" in table.table_html
    assert table.metadata.element_types == ["Table"]
    assert table.metadata.section_title == "Pricing"


def test_every_chunk_carries_metadata(parser):
    content = _docx(
        paragraphs=[("h", "Intro"), ("p", "First section text. " * 40), ("h", "Details"),
                    ("p", "Second section text. " * 40)],
    )
    parsed = _parse(parser, content, DOCX, "report.docx", version=3)

    assert [c.metadata.chunk_index for c in parsed.chunks] == list(range(len(parsed.chunks)))
    sha = hashlib.sha256(content).hexdigest()
    for chunk in parsed.chunks:
        m = chunk.metadata
        assert m.document_version == 3
        assert m.content_sha256 == sha
        assert m.filename == "report.docx"
        assert m.mime_type == DOCX
        assert m.parser_version.startswith("unstructured-")
        assert m.parser_version.endswith(CHUNKING_VERSION)
        # DOCX has no pages.
        assert m.page_start is None and m.page_end is None
        assert m.element_types
    assert parsed.chunks[0].metadata.section_title == "Intro"
    assert parsed.chunks[-1].metadata.section_title == "Details"


@pytest.mark.parametrize(
    ("content", "mime_type"),
    [
        (b"   \n\t ", TEXT),
        (b"", MARKDOWN),
        (b"\xff\xfe\xfa not utf-8", TEXT),
        (b"not a zip file", DOCX),
    ],
)
def test_unreadable_or_empty_files_fail_extraction(parser, content, mime_type):
    with pytest.raises(ExtractionFailed):
        _parse(parser, content, mime_type)


def test_unsupported_type_raises(parser):
    with pytest.raises(UnsupportedFileType):
        _parse(parser, b"\x89PNG", "image/png")


# ---------- PDFs: hi_res, OCR, tables (need tesseract + poppler) ----------


def _text_pdf_with_table() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    styles = getSampleStyleSheet()
    table = Table(ROWS)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4).build(
        [
            Paragraph("Pricing Review", styles["Title"]),
            Paragraph("The subscription cost for next year has been updated. " * 5, styles["Normal"]),
            table,
            PageBreak(),
            Paragraph("Support hours are nine to five on weekdays. " * 8, styles["Normal"]),
        ]
    )
    return out.getvalue()


def _scanned_pdf(text_lines: list[str]) -> bytes:
    """An image-only PDF — no text layer, like a scan."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    for i, line in enumerate(text_lines):
        draw.text((120, 150 + i * 90), line, font=font, fill="black")
    out = io.BytesIO()
    image.save(out, "PDF", resolution=200)
    return out.getvalue()


@needs_ocr
def test_text_pdf_keeps_table_as_its_own_chunk_with_pages(parser):
    parsed = _parse(parser, _text_pdf_with_table(), PDF, "pricing.pdf")

    tables = [c for c in parsed.chunks if c.kind == ChunkKind.TABLE]
    assert tables, "hi_res should detect the table"
    assert "Starter" in tables[0].text and tables[0].table_html
    assert all(c.metadata.page_start is not None for c in parsed.chunks)
    assert max(c.metadata.page_end for c in parsed.chunks) == 2
    assert "Support hours" in parsed.text


@needs_ocr
def test_scanned_pdf_is_read_with_ocr(parser):
    content = _scanned_pdf(["Quarterly Report", "The subscription cost has been updated."])
    parsed = _parse(parser, content, PDF, "scan.pdf")

    assert "subscription" in parsed.text.lower()
    assert parsed.chunks[0].metadata.page_start == 1


@needs_ocr
def test_blank_pdf_fails_even_with_ocr(parser):
    with pytest.raises(ExtractionFailed):
        _parse(parser, _scanned_pdf([]), PDF, "blank.pdf")


def test_spacy_model_is_installed_not_downloaded_at_runtime():
    # unstructured downloads en_core_web_sm into site-packages on first use
    # if it's missing, which fails in the container (non-root user). It's a
    # declared, locked dependency instead (pyproject.toml).
    import spacy

    assert spacy.util.is_package("en_core_web_sm")


@needs_ocr
def test_warm_up_loads_every_model(parser):
    # Runs where the models can be used at all (the Docker image).
    parser.warm_up()
