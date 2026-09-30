import io

import pytest
from docx import Document as DocxDocument

from findr.adapters.outbound.files.file_text_extractor import (
    DOCX,
    MARKDOWN,
    PDF,
    TEXT,
    FileTextExtractor,
    resolve_mime_type,
)
from findr.domain.exceptions import ExtractionFailed, UnsupportedFileType


def _pdf_bytes(text: str | None) -> bytes:
    """A minimal single-page PDF, with a text layer if `text` is given —
    built by hand so the test needs no PDF-writing dependency."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    )
    return out.getvalue()


def _docx_bytes(*paragraphs: str) -> bytes:
    doc = DocxDocument()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell one"
    table.rows[0].cells[1].text = "cell two"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_extracts_plain_text_and_markdown():
    extractor = FileTextExtractor()
    assert extractor.extract(b"hello renewal terms\n", TEXT) == "hello renewal terms"
    assert extractor.extract("# Title\n\ncafé".encode(), MARKDOWN) == "# Title\n\ncafé"


def test_strips_utf8_bom():
    assert FileTextExtractor().extract(b"\xef\xbb\xbfhello", TEXT) == "hello"


def test_extracts_pdf_text_layer():
    assert "Quarterly renewal terms" in FileTextExtractor().extract(
        _pdf_bytes("Quarterly renewal terms"), PDF
    )


def test_extracts_docx_paragraphs_and_tables():
    text = FileTextExtractor().extract(_docx_bytes("First paragraph", "Second one"), DOCX)
    assert "First paragraph" in text
    assert "Second one" in text
    assert "cell one" in text and "cell two" in text


def test_pdf_without_text_layer_is_an_extraction_failure():
    # e.g. a scanned PDF — OCR is out of scope (specs/file-upload.md §1).
    with pytest.raises(ExtractionFailed):
        FileTextExtractor().extract(_pdf_bytes(None), PDF)


def test_whitespace_only_text_is_an_extraction_failure():
    with pytest.raises(ExtractionFailed):
        FileTextExtractor().extract(b"   \n\t ", TEXT)


def test_corrupt_files_are_extraction_failures():
    extractor = FileTextExtractor()
    with pytest.raises(ExtractionFailed):
        extractor.extract(b"not a pdf at all", PDF)
    with pytest.raises(ExtractionFailed):
        extractor.extract(b"not a zip", DOCX)
    with pytest.raises(ExtractionFailed):
        extractor.extract(b"\xff\xfe\xfa invalid utf-8", TEXT)


def test_unsupported_type_raises():
    with pytest.raises(UnsupportedFileType):
        FileTextExtractor().extract(b"\x89PNG", "image/png")


@pytest.mark.parametrize(
    ("filename", "declared", "expected"),
    [
        ("notes.md", "application/octet-stream", MARKDOWN),
        ("notes.md", "", MARKDOWN),
        ("notes.MD", None, MARKDOWN),
        ("report.pdf", "application/pdf", PDF),
        ("letter.txt", "text/plain; charset=utf-8", TEXT),
        ("letter.docx", "application/octet-stream", DOCX),
        ("readme", "text/markdown", MARKDOWN),
        ("photo.png", "image/png", "image/png"),
        ("mystery", None, "application/octet-stream"),
    ],
)
def test_resolve_mime_type(filename, declared, expected):
    assert resolve_mime_type(filename, declared) == expected
