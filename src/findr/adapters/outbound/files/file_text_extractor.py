from __future__ import annotations

import io
from pathlib import PurePath

from findr.domain.exceptions import ExtractionFailed, UnsupportedFileType

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = "text/plain"
MARKDOWN = "text/markdown"

SUPPORTED_MIME_TYPES = frozenset({PDF, DOCX, TEXT, MARKDOWN})

_MIME_BY_EXTENSION = {
    ".pdf": PDF,
    ".docx": DOCX,
    ".txt": TEXT,
    ".md": MARKDOWN,
    ".markdown": MARKDOWN,
}


def resolve_mime_type(filename: str, declared: str | None) -> str:
    """Browsers are inconsistent about Content-Type for uploads — .md in
    particular often arrives as application/octet-stream or empty — so a
    declared type we support wins, and otherwise the extension decides.
    Returns the declared (or octet-stream) type unchanged when neither is
    recognised, so the extractor raises UnsupportedFileType with it."""
    declared = (declared or "").split(";", 1)[0].strip().lower()
    if declared in SUPPORTED_MIME_TYPES:
        return declared
    by_extension = _MIME_BY_EXTENSION.get(PurePath(filename).suffix.lower())
    return by_extension or declared or "application/octet-stream"


class FileTextExtractor:
    """Implements ports.text_extractor.TextExtractor — one branch per file
    type, in one place (same shape as connector_factory.py)."""

    def extract(self, content: bytes, mime_type: str) -> str:
        if mime_type == PDF:
            text = _extract_pdf(content)
        elif mime_type == DOCX:
            text = _extract_docx(content)
        elif mime_type in (TEXT, MARKDOWN):
            text = _extract_plain(content)
        else:
            raise UnsupportedFileType(
                f"Unsupported file type {mime_type!r}; upload a PDF, DOCX, TXT or Markdown file"
            )

        text = text.strip()
        if not text:
            raise ExtractionFailed(
                "No text could be extracted from this file (scanned PDFs aren't supported)"
            )
        return text


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ExtractionFailed(f"Could not read PDF: {exc}") from exc


def _extract_docx(content: bytes) -> str:
    import zipfile

    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(content))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionFailed(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_plain(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionFailed("Text file is not valid UTF-8") from exc
